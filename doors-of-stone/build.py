#!/usr/bin/env python3
"""Assemble chapters/*.md into a Word document (and a PDF via ReportLab).

Chapter file conventions:
  - Optional first line "@@PART <title>" starts a new part (own page).
  - "# <Title>" is the chapter heading. Prefixes "Chapter ...:" are stripped;
    titles beginning with "Interlude", "PROLOGUE" or "EPILOGUE" are not numbered.
  - Paragraphs are separated by blank lines; "---" is a scene break.
  - *text* renders as italic.
"""
import glob
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(HERE, "output")
TITLE = "The Doors of Stone"
SUBTITLE = "Day Three of the Kingkiller Chronicle"
NOTE = (
    "An unofficial fan continuation, written in homage to Patrick Rothfuss's "
    "The Name of the Wind and The Wise Man's Fear. It is not written, authorized, "
    "or endorsed by Patrick Rothfuss or his publishers, and it is not the "
    "forthcoming novel of the same name. The characters and world of the "
    "Kingkiller Chronicle belong to their creator."
)

NUMBER_WORDS = [
    "Zero", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine",
    "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen",
    "Seventeen", "Eighteen", "Nineteen",
]
TENS = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]


def number_to_words(n):
    if n < 20:
        return NUMBER_WORDS[n]
    if n < 100:
        t, o = divmod(n, 10)
        return TENS[t] + ("-" + NUMBER_WORDS[o] if o else "")
    h, r = divmod(n, 100)
    return NUMBER_WORDS[h] + " Hundred" + (" " + number_to_words(r) if r else "")


def add_runs(paragraph, text):
    parts = re.split(r"(\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def set_base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    normal.font.size = Pt(11.5)
    pf = normal.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.25
    pf.first_line_indent = Inches(0.3)
    for sec in doc.sections:
        sec.page_width = Inches(6)
        sec.page_height = Inches(9)
        sec.left_margin = sec.right_margin = Inches(0.8)
        sec.top_margin = sec.bottom_margin = Inches(0.8)


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def centered(doc, text, size, bold=False, italic=False, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return p


def parse_chapter(path):
    with open(path, encoding="utf-8") as f:
        lines = f.read().split("\n")
    part = None
    if lines and lines[0].startswith("@@PART"):
        part = lines[0][len("@@PART"):].strip()
        lines = lines[1:]
    while lines and not lines[0].strip():
        lines = lines[1:]
    title = lines[0].lstrip("#").strip() if lines and lines[0].startswith("#") else ""
    body = "\n".join(lines[1:]).strip()
    title = re.sub(r"^Chapter [A-Za-z\- ]+:\s*", "", title)
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()]
    return part, title, paragraphs


def main():
    files = sorted(glob.glob(os.path.join(HERE, "chapters", "*.md")))
    doc = Document()
    set_base_styles(doc)
    add_page_number_footer(doc.sections[0])

    centered(doc, TITLE.upper(), 30, bold=True, space_before=150)
    centered(doc, SUBTITLE, 14, italic=True, space_before=12)
    centered(doc, "A continuation in the spirit of Patrick Rothfuss", 11, italic=True, space_before=40)
    page_break(doc)
    centered(doc, "A Note", 13, bold=True, space_before=80, space_after=14)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, NOTE)

    chapter_no = 0
    words = 0
    for path in files:
        part, title, paragraphs = parse_chapter(path)
        if part:
            page_break(doc)
            centered(doc, part, 22, bold=True, space_before=200)
        page_break(doc)
        upper = title.upper()
        numbered = not (upper.startswith("INTERLUDE") or upper.startswith("PROLOGUE")
                        or upper.startswith("EPILOGUE"))
        if numbered:
            chapter_no += 1
            centered(doc, "Chapter " + number_to_words(chapter_no), 12, space_before=90)
            centered(doc, title, 18, bold=True, space_before=6, space_after=30)
        else:
            centered(doc, title, 18, bold=True, space_before=90, space_after=30)
        first = True
        for para in paragraphs:
            if para == "---":
                centered(doc, "*   *   *", 11, space_before=8, space_after=8)
                first = True
                continue
            p = doc.add_paragraph()
            if first:
                p.paragraph_format.first_line_indent = Inches(0)
                first = False
            text = para.replace("\n", " ")
            words += len(text.split())
            add_runs(p, text)

    os.makedirs(OUT_DIR, exist_ok=True)
    docx_path = os.path.join(OUT_DIR, "The_Doors_of_Stone.docx")
    doc.save(docx_path)
    print(f"chapters numbered: {chapter_no}; files: {len(files)}; words: {words}")
    print("wrote", docx_path)
    if "--pdf" in sys.argv:
        pdf_path = os.path.join(OUT_DIR, "The_Doors_of_Stone.pdf")
        build_pdf(files, pdf_path)
        print("wrote", pdf_path)

def build_pdf(files, pdf_path):
    from xml.sax.saxutils import escape

    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (Flowable, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer)

    font_dir = "/usr/share/fonts/truetype/liberation"
    for name, fname in (("Serif", "LiberationSerif-Regular.ttf"),
                        ("Serif-Bold", "LiberationSerif-Bold.ttf"),
                        ("Serif-Italic", "LiberationSerif-Italic.ttf"),
                        ("Serif-BoldItalic", "LiberationSerif-BoldItalic.ttf")):
        pdfmetrics.registerFont(TTFont(name, os.path.join(font_dir, fname)))
    pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-Bold",
                                  italic="Serif-Italic", boldItalic="Serif-BoldItalic")

    body = ParagraphStyle("body", fontName="Serif", fontSize=11, leading=15,
                          alignment=TA_JUSTIFY, firstLineIndent=0.3 * inch)
    body_first = ParagraphStyle("body_first", parent=body, firstLineIndent=0)
    center = ParagraphStyle("center", parent=body, alignment=TA_CENTER, firstLineIndent=0)
    chap_no = ParagraphStyle("chap_no", parent=center, fontSize=12, leading=16)
    chap_title = ParagraphStyle("chap_title", parent=center, fontName="Serif-Bold",
                                fontSize=18, leading=24)
    part_style = ParagraphStyle("part", parent=center, fontName="Serif-Bold",
                                fontSize=22, leading=30)
    title_style = ParagraphStyle("title", parent=center, fontName="Serif-Bold",
                                 fontSize=28, leading=36)
    sub_style = ParagraphStyle("sub", parent=center, fontName="Serif-Italic",
                               fontSize=13, leading=18)

    def markup(text):
        text = escape(text)
        return re.sub(r"\*([^*]+)\*", r"<i>\1</i>", text)

    class Outline(Flowable):
        def __init__(self, title, key, level=0):
            super().__init__()
            self.title, self.key, self.level = title, key, level
            self.width = self.height = 0

        def draw(self):
            self.canv.bookmarkPage(self.key)
            self.canv.addOutlineEntry(self.title, self.key, level=self.level)

    def footer(canvas, doc):
        if doc.page > 2:
            canvas.saveState()
            canvas.setFont("Serif", 9)
            canvas.drawCentredString(3 * inch, 0.45 * inch, str(doc.page))
            canvas.restoreState()

    story = [Spacer(1, 2.2 * inch), Paragraph(TITLE.upper(), title_style),
             Spacer(1, 0.2 * inch), Paragraph(SUBTITLE, sub_style),
             Spacer(1, 0.6 * inch),
             Paragraph("A continuation in the spirit of Patrick Rothfuss", sub_style),
             PageBreak(), Spacer(1, 1.2 * inch),
             Paragraph("<b>A Note</b>", center), Spacer(1, 0.2 * inch),
             Paragraph(markup(NOTE), center)]

    chapter_no = 0
    part_level_open = False
    for i, path in enumerate(files):
        part, title, paragraphs = parse_chapter(path)
        if part:
            story += [PageBreak(), Outline(part, "part%d" % i, 0),
                      Spacer(1, 2.8 * inch), Paragraph(markup(part), part_style)]
            part_level_open = True
        story.append(PageBreak())
        upper = title.upper()
        numbered = not (upper.startswith("INTERLUDE") or upper.startswith("PROLOGUE")
                        or upper.startswith("EPILOGUE"))
        level = 1 if part_level_open else 0
        if numbered:
            chapter_no += 1
            label = "Chapter " + number_to_words(chapter_no)
            story += [Outline("%s: %s" % (label, title), "ch%d" % i, level),
                      Spacer(1, 1.1 * inch), Paragraph(label, chap_no), Spacer(1, 6)]
        else:
            story += [Outline(title, "ch%d" % i, level), Spacer(1, 1.2 * inch)]
        story += [Paragraph(markup(title), chap_title), Spacer(1, 0.4 * inch)]
        first = True
        for para in paragraphs:
            if para == "---":
                story += [Spacer(1, 6), Paragraph("*&nbsp;&nbsp;&nbsp;*&nbsp;&nbsp;&nbsp;*", center),
                          Spacer(1, 6)]
                first = True
                continue
            style = body_first if first else body
            first = False
            story.append(Paragraph(markup(para.replace("\n", " ")), style))

    doc = SimpleDocTemplate(pdf_path, pagesize=(6 * inch, 9 * inch),
                            leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                            topMargin=0.75 * inch, bottomMargin=0.8 * inch,
                            title=TITLE, author="Unofficial fan continuation",
                            subject=SUBTITLE)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    main()
