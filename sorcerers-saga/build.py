#!/usr/bin/env python3
"""Build KDP 6 x 9 in paperback interiors (PDF) and Word files for the Lattice Chronicles.

Usage:
    python3 build.py            # both books
    python3 build.py book1      # one book
    python3 build.py book1 --pdf-only

Chapter file conventions (book*/chapters/*.md, built in filename order):
  - Optional first line "@@PART <title>" starts a new part on its own right-hand page.
  - "# <Title>" is the chapter heading. Titles starting with "Prologue", "Epilogue" or
    "Interlude" are not numbered.
  - Blank lines separate paragraphs. "---" is a scene break.
  - *italic* and **bold** inline.
  - Lines starting with "> " are verse (centered italic).
  - ":::system", ":::root" and ":::error" open a System panel; ":::" closes it.
"""
import glob
import os
import re
import sys

from reportlab.lib.colors import Color, black
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Flowable, Frame, KeepTogether, PageBreak, PageTemplate,
                                Paragraph, Spacer, Table, TableStyle)
from reportlab.platypus.doctemplate import BaseDocTemplate
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.lib.fonts import addMapping

HERE = os.path.dirname(os.path.abspath(__file__))
FONTS = os.path.join(HERE, "fonts")
OUT = os.path.join(HERE, "output")

SERIES = "The Lattice Chronicles"
AUTHOR = "VerdantLife LLC"
BOOKS = {
    "book1": {
        "title": "The Sorcerer's Reckoning",
        "number": "Book One",
        "file": "The-Sorcerers-Reckoning",
        "dedication": "For every kid who was told what they would be\nbefore anyone asked who they were.",
        "epigraph": ("Little light, little light, burning low,\nkeep the door for the ones who go.",
                     "Ashbourne lullaby"),
    },
    "book2": {
        "title": "The Sorcerer's Legacy",
        "number": "Book Two",
        "file": "The-Sorcerers-Legacy",
        "dedication": "For fathers who come home late,\nand for the ones who leave the light on.",
        "epigraph": ("If the dark comes calling, don't you mind—\na lantern's lit for the ones behind.",
                     "Ashbourne lullaby"),
    },
}

# ---------------------------------------------------------------- fonts
def register_fonts():
    f = lambda n: os.path.join(FONTS, n)
    pdfmetrics.registerFont(TTFont("Garamond", f("EBGaramond-400-normal.ttf")))
    pdfmetrics.registerFont(TTFont("Garamond-Italic", f("EBGaramond-400-italic.ttf")))
    pdfmetrics.registerFont(TTFont("Garamond-Bold", f("EBGaramond-700-normal.ttf")))
    pdfmetrics.registerFont(TTFont("Garamond-BoldItalic", f("EBGaramond-700-italic.ttf")))
    pdfmetrics.registerFont(TTFont("Garamond-Semi", f("EBGaramond-600-normal.ttf")))
    addMapping("Garamond", 0, 0, "Garamond")
    addMapping("Garamond", 0, 1, "Garamond-Italic")
    addMapping("Garamond", 1, 0, "Garamond-Bold")
    addMapping("Garamond", 1, 1, "Garamond-BoldItalic")
    pdfmetrics.registerFont(TTFont("Cinzel", f("Cinzel-400.ttf")))
    pdfmetrics.registerFont(TTFont("Cinzel-Bold", f("Cinzel-700.ttf")))
    pdfmetrics.registerFont(TTFont("Mono", f("PlexMono-400.ttf")))
    pdfmetrics.registerFont(TTFont("Mono-Bold", f("PlexMono-600.ttf")))
    pdfmetrics.registerFont(TTFont("Mono-Italic", f("PlexMono-400i.ttf")))
    addMapping("Mono", 0, 0, "Mono")
    addMapping("Mono", 1, 0, "Mono-Bold")
    addMapping("Mono", 0, 1, "Mono-Italic")
    addMapping("Mono", 1, 1, "Mono-Bold")
    # Fallbacks for glyphs the subsetted web fonts lack (block shading, arrows,
    # geometric shapes, combining marks). DejaVu is freely embeddable.
    pdfmetrics.registerFont(TTFont("MonoFB", f("DejaVuSansMono.ttf")))
    pdfmetrics.registerFont(TTFont("MonoFB-Bold", f("DejaVuSansMono-Bold.ttf")))
    pdfmetrics.registerFont(TTFont("SansFB", f("DejaVuSans.ttf")))
    addMapping("MonoFB", 0, 0, "MonoFB")
    addMapping("MonoFB", 1, 0, "MonoFB-Bold")
    addMapping("MonoFB", 0, 1, "MonoFB")
    addMapping("MonoFB", 1, 1, "MonoFB-Bold")
    from fontTools.ttLib import TTFont as _FT
    CMAPS["mono"] = set(_FT(f("PlexMono-400.ttf")).getBestCmap())
    CMAPS["body"] = (set(_FT(f("EBGaramond-400-normal.ttf")).getBestCmap())
                     & set(_FT(f("EBGaramond-400-italic.ttf")).getBestCmap()))


CMAPS = {}


def with_fallback(text, which, fb):
    """Wrap runs of characters missing from the primary font in <font name=fb>."""
    cmap = CMAPS.get(which)
    if not cmap:
        return text
    out, run = [], []
    for ch in text:
        if ord(ch) > 127 and ord(ch) not in cmap:
            run.append(ch)
        else:
            if run:
                out.append('<font name="%s">%s</font>' % (fb, "".join(run)))
                run = []
            out.append(ch)
    if run:
        out.append('<font name="%s">%s</font>' % (fb, "".join(run)))
    return "".join(out)


# ---------------------------------------------------------------- page geometry
PAGE_W, PAGE_H = 6 * inch, 9 * inch
INSIDE, OUTSIDE = 0.85 * inch, 0.6 * inch
TOP, BOTTOM = 0.8 * inch, 0.75 * inch
BODY_SIZE, BODY_LEAD = 11.5, 15.2

GREY = Color(0.93, 0.93, 0.93)
DARK = Color(0.25, 0.25, 0.25)

# ---------------------------------------------------------------- styles
def styles():
    s = {}
    s["body"] = ParagraphStyle("body", fontName="Garamond", fontSize=BODY_SIZE, leading=BODY_LEAD,
                               alignment=TA_JUSTIFY, firstLineIndent=0.22 * inch, allowWidows=0,
                               allowOrphans=0, hyphenationLang=None)
    s["first"] = ParagraphStyle("first", parent=s["body"], firstLineIndent=0)
    s["verse"] = ParagraphStyle("verse", parent=s["body"], fontName="Garamond-Italic",
                                alignment=TA_CENTER, firstLineIndent=0, spaceBefore=2, spaceAfter=2)
    s["break"] = ParagraphStyle("break", parent=s["body"], alignment=TA_CENTER, firstLineIndent=0,
                                spaceBefore=8, spaceAfter=8)
    s["chnum"] = ParagraphStyle("chnum", fontName="Cinzel", fontSize=12, leading=16,
                                alignment=TA_CENTER, textColor=DARK)
    s["chtitle"] = ParagraphStyle("chtitle", fontName="Cinzel-Bold", fontSize=20, leading=26,
                                  alignment=TA_CENTER, spaceBefore=6)
    s["parttitle"] = ParagraphStyle("parttitle", fontName="Cinzel-Bold", fontSize=24, leading=32,
                                    alignment=TA_CENTER)
    s["partnum"] = ParagraphStyle("partnum", fontName="Cinzel", fontSize=14, leading=20,
                                  alignment=TA_CENTER, textColor=DARK)
    s["sys"] = ParagraphStyle("sys", fontName="Mono", fontSize=8.3, leading=10.6, alignment=TA_LEFT)
    s["center"] = ParagraphStyle("center", fontName="Garamond", fontSize=11, leading=15,
                                 alignment=TA_CENTER)
    s["small"] = ParagraphStyle("small", fontName="Garamond", fontSize=9, leading=12,
                                alignment=TA_LEFT)
    s["toc0"] = ParagraphStyle("toc0", fontName="Cinzel-Bold", fontSize=10.5, leading=16,
                               spaceBefore=8)
    s["toc1"] = ParagraphStyle("toc1", fontName="Garamond", fontSize=11, leading=14.5,
                               leftIndent=12)
    return s


# ---------------------------------------------------------------- markers
class Marker(Flowable):
    """Zero-size flowable intercepted by the doc template."""
    def __init__(self, kind, value=None):
        super().__init__()
        self.kind, self.value = kind, value

    def wrap(self, *a):
        return 0, 0

    def draw(self):
        pass


class Ornament(Flowable):
    """A small diamond rule, drawn as vectors so it never depends on font coverage."""
    def __init__(self, width=60, gap=10):
        super().__init__()
        self.width, self.gap = width, gap

    def wrap(self, aw, ah):
        self.aw = aw
        return aw, self.gap * 2

    def draw(self):
        c = self.canv
        cx, cy = self.aw / 2, self.gap
        c.setLineWidth(0.6)
        c.line(cx - self.width, cy, cx - 6, cy)
        c.line(cx + 6, cy, cx + self.width, cy)
        p = c.beginPath()
        p.moveTo(cx, cy + 3.5); p.lineTo(cx + 3.5, cy); p.lineTo(cx, cy - 3.5); p.lineTo(cx - 3.5, cy); p.close()
        c.drawPath(p, fill=1, stroke=0)


class SceneBreak(Flowable):
    def wrap(self, aw, ah):
        self.aw = aw
        return aw, 22

    def draw(self):
        c = self.canv
        y = 11
        for dx in (-14, 0, 14):
            x = self.aw / 2 + dx
            p = c.beginPath()
            p.moveTo(x, y + 2.6); p.lineTo(x + 2.6, y); p.lineTo(x, y - 2.6); p.lineTo(x - 2.6, y); p.close()
            c.drawPath(p, fill=1, stroke=0)


# ---------------------------------------------------------------- doc template
class BookDoc(BaseDocTemplate):
    def __init__(self, path, book, **kw):
        super().__init__(path, pagesize=(PAGE_W, PAGE_H), leftMargin=INSIDE, rightMargin=OUTSIDE,
                         topMargin=TOP, bottomMargin=BOTTOM, title=book["title"], author=AUTHOR,
                         subject=SERIES, **kw)
        self.book = book
        fw = PAGE_W - INSIDE - OUTSIDE
        fh = PAGE_H - TOP - BOTTOM
        recto = Frame(INSIDE, BOTTOM, fw, fh, id="recto", leftPadding=0, rightPadding=0,
                      topPadding=0, bottomPadding=0)
        verso = Frame(OUTSIDE, BOTTOM, fw, fh, id="verso", leftPadding=0, rightPadding=0,
                      topPadding=0, bottomPadding=0)
        self.addPageTemplates([
            PageTemplate("recto", [recto], onPageEnd=self.decorate),
            PageTemplate("verso", [verso], onPageEnd=self.decorate),
        ])
        self.reset_state()

    def reset_state(self):
        self.blank_pages, self.opener_pages = set(), set()
        self.body_started = False
        self.body_start_page = None
        self.current_title = ""
        self.toc_key = 0

    def beforeDocument(self):
        self.reset_state()

    def afterPage(self):
        nxt = self.page + 1
        self._nextPageTemplateIndex = 0 if nxt % 2 == 1 else 1

    def handle_flowable(self, flowables):
        f = flowables[0]
        if isinstance(f, Marker):
            flowables.pop(0)
            empty = self._curPageFlowableCount == 0
            if f.kind == "recto":
                if empty:
                    ins = [] if self.page % 2 == 1 else [Marker("blank"), PageBreak()]
                else:
                    ins = [PageBreak(), Marker("blank"), PageBreak()] if self.page % 2 == 1 else [PageBreak()]
                flowables[0:0] = ins
            elif f.kind == "newpage":
                if not empty:
                    flowables.insert(0, PageBreak())
            elif f.kind == "blank":
                self.blank_pages.add(self.page)
                flowables.insert(0, Spacer(1, 1))
            elif f.kind == "opener":
                self.opener_pages.add(self.page)
                if f.value is not None:
                    self.current_title = f.value
            elif f.kind == "body":
                self.body_started = True
                self.body_start_page = self.page
            elif f.kind == "toc":
                level, text = f.value
                self.toc_key += 1
                key = "k%d" % self.toc_key
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(re.sub(r"<[^>]+>", "", text), key, level=0, closed=True)
                self.notify("TOCEntry", (level, text, self.page, key))
            return
        super().handle_flowable(flowables)

    def decorate(self, canv, doc):
        p = doc.page
        if not self.body_started or p in self.blank_pages:
            return
        recto = p % 2 == 1
        canv.saveState()
        if p not in self.opener_pages:
            canv.setFont("Cinzel", 8)
            text = (self.current_title if recto else self.book["title"]).upper()
            y = PAGE_H - TOP + 0.32 * inch
            canv.drawCentredString((INSIDE if recto else OUTSIDE) + (PAGE_W - INSIDE - OUTSIDE) / 2, y, text)
            canv.setFont("Garamond", 10)
            if recto:
                canv.drawRightString(PAGE_W - OUTSIDE, y, str(p))
            else:
                canv.drawString(OUTSIDE, y, str(p))
        else:
            canv.setFont("Garamond", 10)
            canv.drawCentredString((INSIDE if recto else OUTSIDE) + (PAGE_W - INSIDE - OUTSIDE) / 2,
                                   BOTTOM - 0.38 * inch, str(p))
        canv.restoreState()


# ---------------------------------------------------------------- markdown
def inline(text):
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"<i>\1</i>", text)
    return with_fallback(text, "body", "SansFB")


def parse_chapter(path):
    with open(path, encoding="utf-8") as fh:
        lines = fh.read().split("\n")
    part = None
    if lines and lines[0].startswith("@@PART"):
        part = lines[0][len("@@PART"):].strip()
        lines = lines[1:]
    while lines and not lines[0].strip():
        lines.pop(0)
    title = lines.pop(0).lstrip("#").strip() if lines and lines[0].startswith("#") else ""
    blocks, para, i = [], [], 0

    def flush():
        if para:
            blocks.append(("p", " ".join(x.strip() for x in para)))
            para.clear()

    while i < len(lines):
        ln = lines[i]
        st = ln.strip()
        if st.startswith(":::") and len(st) > 3:
            flush()
            kind = st[3:].strip()
            body = []
            i += 1
            while i < len(lines) and lines[i].strip() != ":::":
                body.append(lines[i].rstrip())
                i += 1
            blocks.append(("sys", kind, body))
        elif st == "---":
            flush()
            blocks.append(("break",))
        elif st.startswith(">"):
            flush()
            verse = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                verse.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(("verse", verse))
            continue
        elif not st:
            flush()
        else:
            para.append(ln)
        i += 1
    flush()
    return part, title, blocks


def system_panel(kind, lines, S, width):
    def fmt(l):
        t = l.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace(" ", "&nbsp;")
        return with_fallback(t, "mono", "MonoFB") or "&nbsp;"
    rows = []
    for idx, l in enumerate(lines):
        txt = fmt(l)
        if idx == 0:
            txt = "<b>%s</b>" % txt
        if kind == "error":
            txt = "<i>%s</i>" % txt
        rows.append([Paragraph(txt, S["sys"])])
    t = Table(rows, colWidths=[width - 0.3 * inch])
    style = [
        ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, 0), 6), ("BOTTOMPADDING", (0, -1), (-1, -1), 6),
    ]
    if kind == "root":
        style += [("BOX", (0, 0), (-1, -1), 1.6, black), ("INNERGRID", (0, 0), (-1, -1), 0, GREY),
                  ("BACKGROUND", (0, 0), (-1, -1), Color(0.97, 0.97, 0.97))]
    elif kind == "error":
        style += [("BOX", (0, 0), (-1, -1), 0.8, black, None, (2, 2)),
                  ("BACKGROUND", (0, 0), (-1, -1), Color(0.9, 0.9, 0.9))]
    else:
        style += [("BOX", (0, 0), (-1, -1), 0.6, black), ("BACKGROUND", (0, 0), (-1, -1), GREY)]
    t.setStyle(TableStyle(style))
    t.hAlign = "CENTER"
    return KeepTogether([Spacer(1, 5), t, Spacer(1, 7)]) if len(lines) < 40 else t


# ---------------------------------------------------------------- assembly
def front_matter(book, S, toc):
    fl = []
    fl += [Spacer(1, 2.2 * inch), Paragraph(book["title"].upper(), S["parttitle"])]
    fl += [Marker("newpage")]
    fl += [PageBreak(), Marker("blank"), PageBreak()]
    # title page (recto)
    fl += [Spacer(1, 1.6 * inch), Paragraph(book["title"].upper(), ParagraphStyle(
        "tt", parent=S["parttitle"], fontSize=28, leading=36))]
    fl += [Spacer(1, 10), Ornament(70), Spacer(1, 10)]
    fl += [Paragraph("%s of %s" % (book["number"], SERIES), ParagraphStyle(
        "ts", parent=S["center"], fontName="Garamond-Italic", fontSize=14, leading=18))]
    fl += [Spacer(1, 2.6 * inch), Paragraph(AUTHOR.upper(), S["partnum"])]
    fl += [PageBreak()]
    # copyright (verso)
    cp = [
        "%s<br/>%s of %s" % (book["title"], book["number"], SERIES),
        "Copyright © 2026 %s. All rights reserved." % AUTHOR,
        "No part of this book may be reproduced in any form without written permission from the publisher, except for brief quotations in reviews.",
        "This is a work of fiction. Names, characters, places and events are products of the author's imagination. Any resemblance to actual persons, living or dead, is coincidental.",
        "First edition.",
    ]
    fl += [Spacer(1, 4.2 * inch)] + [Paragraph(x, ParagraphStyle("cp", parent=S["small"], spaceAfter=7)) for x in cp]
    fl += [PageBreak()]
    # dedication (recto)
    fl += [Spacer(1, 2.4 * inch), Paragraph(inline(book["dedication"]).replace("\n", "<br/>"),
                                            ParagraphStyle("ded", parent=S["center"], fontName="Garamond-Italic", fontSize=12, leading=17))]
    fl += [PageBreak(), Marker("blank"), PageBreak()]
    # contents (recto)
    fl += [Paragraph("CONTENTS", S["chtitle"]), Spacer(1, 6), Ornament(40), Spacer(1, 10), toc]
    fl += [Marker("recto")]
    # epigraph
    ep, src = book["epigraph"]
    fl += [Spacer(1, 2.6 * inch), Paragraph(inline(ep).replace("\n", "<br/>"),
                                            ParagraphStyle("ep", parent=S["verse"], fontSize=12, leading=17)),
           Spacer(1, 8), Paragraph("— " + src, ParagraphStyle("eps", parent=S["center"], fontSize=10))]
    return fl


def chapter_flowables(path, num, S, width):
    part, title, blocks = parse_chapter(path)
    fl = []
    if part:
        m = re.match(r"(Part [A-Za-z]+)\s*:\s*(.+)", part)
        pnum, ptitle = (m.group(1), m.group(2)) if m else ("", part)
        fl += [Marker("recto"), Marker("opener", None), Marker("toc", (0, part)),
               Spacer(1, 2.4 * inch), Paragraph(pnum.upper(), S["partnum"]), Spacer(1, 10),
               Ornament(60), Spacer(1, 12), Paragraph(ptitle.upper(), S["parttitle"]),
               Marker("recto")]
    unnumbered = re.match(r"(Prologue|Epilogue|Interlude)", title, re.I)
    label = ""
    if unnumbered:
        if ":" in title:
            label, title = [x.strip() for x in title.split(":", 1)]
        else:
            label, title = title, ""
    else:
        num += 1
        label = "Chapter %d" % num
    fl += [Marker("newpage"), Marker("opener", title or label),
           Marker("toc", (1, ("%s: %s" % (label, title)) if title else label)),
           Spacer(1, 1.1 * inch), Paragraph(label.upper(), S["chnum"])]
    if title:
        fl += [Paragraph(inline(title), S["chtitle"])]
    fl += [Spacer(1, 6), Ornament(45), Spacer(1, 22)]
    first = True
    words = 0
    for b in blocks:
        if b[0] == "p":
            words += len(b[1].split())
            txt = inline(b[1])
            if first:
                # small caps style lead-in: uppercase the first few words
                m = re.match(r"((?:\S+\s+){0,3}\S+)(.*)", txt, re.S)
                lead, rest = m.group(1), m.group(2)
                if "<" not in lead:
                    txt = '<font name="Garamond-Semi">%s</font>%s' % (lead.upper(), rest)
                fl.append(Paragraph(txt, S["first"]))
                first = False
            else:
                fl.append(Paragraph(txt, S["body"]))
        elif b[0] == "break":
            fl.append(SceneBreak())
            first = True
        elif b[0] == "verse":
            fl.append(Spacer(1, 4))
            for v in b[1]:
                fl.append(Paragraph(inline(v), S["verse"]))
            fl.append(Spacer(1, 6))
            words += sum(len(v.split()) for v in b[1])
        elif b[0] == "sys":
            fl.append(system_panel(b[1], b[2], S, width))
            words += sum(len(x.split()) for x in b[2])
    return fl, num, words


def build_pdf(key):
    book = BOOKS[key]
    S = styles()
    width = PAGE_W - INSIDE - OUTSIDE
    files = sorted(glob.glob(os.path.join(HERE, key, "chapters", "*.md")))
    toc = TableOfContents(levelStyles=[S["toc0"], S["toc1"]], dotsMinLevel=1)
    story = front_matter(book, S, toc)
    story.append(Marker("body"))
    num, total = 0, 0
    for path in files:
        fl, num, words = chapter_flowables(path, num, S, width)
        story += fl
        total += words
    # back matter
    story += [Marker("recto"), Marker("opener", ""), Spacer(1, 2.6 * inch),
              Paragraph("THE STORY CONTINUES", S["chnum"]), Spacer(1, 8), Ornament(50), Spacer(1, 10),
              Paragraph(("Book Two of %s<br/><i>The Sorcerer's Legacy</i>" if key == "book1" else
                         "Book Three of %s<br/><i>coming soon</i>") % SERIES, S["center"])]
    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, "%s-KDP-6x9-Interior.pdf" % book["file"])
    doc = BookDoc(path, book)
    doc.multiBuild(story)
    pages = doc.page
    print("%s: %d words, %d pages -> %s" % (book["title"], total, pages, os.path.relpath(path, HERE)))
    return total, pages


# ---------------------------------------------------------------- docx
def build_docx(key):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    book = BOOKS[key]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(6), Inches(9)
    sec.left_margin, sec.right_margin = Inches(0.85), Inches(0.6)
    sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.75)
    sec.gutter = Inches(0)
    settings = doc.settings.element
    mirror = OxmlElement("w:mirrorMargins")
    settings.append(mirror)

    normal = doc.styles["Normal"]
    normal.font.name = "EB Garamond"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "EB Garamond")
    normal.font.size = Pt(11.5)
    pf = normal.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = Pt(15.2)
    pf.first_line_indent = Inches(0.22)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True

    def runs(p, text):
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                p.add_run(part[2:-2]).bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                p.add_run(part[1:-1]).italic = True
            else:
                p.add_run(part)

    def center(text, size, bold=False, italic=False, before=0, after=0, font=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold, r.italic = bold, italic
        if font:
            r.font.name = font
        return p

    def page_break():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    center(book["title"].upper(), 26, bold=True, before=140, font="Cinzel")
    center("%s of %s" % (book["number"], SERIES), 13, italic=True, before=12)
    center(AUTHOR, 12, before=160)
    page_break()
    center("Copyright © 2026 %s. All rights reserved." % AUTHOR, 9, before=300)
    center("This is a work of fiction.", 9)
    page_break()
    center(book["dedication"], 12, italic=True, before=180)
    page_break()

    num = 0
    for path in sorted(glob.glob(os.path.join(HERE, key, "chapters", "*.md"))):
        part, title, blocks = parse_chapter(path)
        if part:
            m = re.match(r"(Part [A-Za-z]+)\s*:\s*(.+)", part)
            center(m.group(1).upper() if m else "", 14, before=180, font="Cinzel")
            center((m.group(2) if m else part).upper(), 22, bold=True, before=8, font="Cinzel")
            page_break()
        if re.match(r"(Prologue|Epilogue|Interlude)", title, re.I):
            label, t = ([x.strip() for x in title.split(":", 1)] + [""])[:2] if ":" in title else (title, "")
        else:
            num += 1
            label, t = "Chapter %d" % num, title
        center(label.upper(), 12, before=70, font="Cinzel")
        if t:
            center(t, 19, bold=True, before=4, after=24, font="Cinzel")
        first = True
        for b in blocks:
            if b[0] == "p":
                p = doc.add_paragraph()
                if first:
                    p.paragraph_format.first_line_indent = Inches(0)
                    first = False
                runs(p, b[1])
            elif b[0] == "break":
                center("◆   ◆   ◆", 9, before=6, after=6)
                first = True
            elif b[0] == "verse":
                for v in b[1]:
                    p = center("", 11.5)
                    runs(p, v)
                    for r in p.runs:
                        r.italic = True
            elif b[0] == "sys":
                tbl = doc.add_table(rows=1, cols=1)
                tbl.style = "Table Grid"
                cell = tbl.rows[0].cells[0]
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "EEEEEE")
                cell._tc.get_or_add_tcPr().append(shd)
                cell.paragraphs[0].text = ""
                for idx, l in enumerate(b[2]):
                    p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
                    p.paragraph_format.first_line_indent = Inches(0)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    r = p.add_run(l if l else " ")
                    r.font.name = "Courier New"
                    r.font.size = Pt(8)
                    r.bold = idx == 0
                    r.italic = b[1] == "error"
                doc.add_paragraph()
        page_break()
    path = os.path.join(OUT, "%s-KDP-6x9.docx" % book["file"])
    doc.save(path)
    print("Word file -> %s" % os.path.relpath(path, HERE))


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    keys = args or ["book1", "book2"]
    register_fonts()
    for k in keys:
        if not glob.glob(os.path.join(HERE, k, "chapters", "*.md")):
            print("%s: no chapters yet" % k)
            continue
        build_pdf(k)
        if "--pdf-only" not in sys.argv:
            build_docx(k)


if __name__ == "__main__":
    main()
